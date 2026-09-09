from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(r"C:\Users\jadgh\Desktop\MSFEA-Chatbot\MSFEA-Chatbot\output\documents\MSFEA_LLM_Options_Comparison.docx")

BLACK = "111111"
DARK = "30343B"
MID = "5B6470"
LIGHT = "F2F4F6"
VERY_LIGHT = "F8F9FA"
BORDER = "C9CED4"
ACCENT = "8A2432"


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 70, start: int = 80, bottom: int = 70, end: int = 80) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color: str = BORDER, size: str = "5") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_width(cell, width_inches: float) -> None:
    cell.width = Inches(width_inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_keep_with_next(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    keep = OxmlElement("w:keepNext")
    p_pr.append(keep)


def set_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    tr_pr.append(cant)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), ACCENT)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_props.append(color)
    run_props.append(underline)
    run.append(run_props)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_1, instr_text, fld_char_2])
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MID)


def style_paragraph(paragraph, *, before: float = 0, after: float = 4, line: float = 1.04) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def add_heading(doc: Document, text: str, level: int = 1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.add_run(text)
    set_keep_with_next(paragraph)
    return paragraph


def add_bullet(doc: Document, lead: str, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(lead)
    r.bold = True
    p.add_run(text)
    style_paragraph(p, after=2.5, line=1.02)


def add_numbered(doc: Document, lead: str, text: str) -> None:
    p = doc.add_paragraph(style="List Number")
    r = p.add_run(lead)
    r.bold = True
    p.add_run(text)
    style_paragraph(p, after=2.5, line=1.02)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.25)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.4)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.04

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(9.1)
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.left_indent = Inches(0.2)
        style.paragraph_format.first_line_indent = Inches(-0.12)

    h1 = styles["Heading 1"]
    h1.font.name = "Arial"
    h1.font.size = Pt(13.5)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor.from_string(BLACK)
    h1.paragraph_format.space_before = Pt(7)
    h1.paragraph_format.space_after = Pt(3)
    h1.paragraph_format.keep_with_next = True

    h2 = styles["Heading 2"]
    h2.font.name = "Arial"
    h2.font.size = Pt(11.2)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor.from_string(BLACK)
    h2.paragraph_format.space_before = Pt(5)
    h2.paragraph_format.space_after = Pt(2)
    h2.paragraph_format.keep_with_next = True

    for sec in doc.sections:
        footer_p = sec.footer.paragraphs[0]
        add_page_number(footer_p)


def add_comparison_table(doc: Document) -> None:
    headers = ["Criterion", "Gemini Flash-Lite (current)", "Ollama local: Llama 3.1 8B", "Ollama Cloud"]
    rows = [
        (
            "Answer quality",
            "Best result in this project. Live tests gave direct, grounded answers and correct rule application.",
            "In my trial, answers were weaker. A small 8B model can work, but needs prompt retuning and a full golden-set run.",
            "Depends on the chosen cloud model. It must pass the same evaluation before use.",
        ),
        (
            "Response time",
            "Measured provider time: 0.86-1.95 seconds across five live reasoning cases.",
            "Hardware-dependent. CPU-only hosting is expected to be much slower; a GPU is the practical route.",
            "Managed datacenter compute; likely responsive, but not measured in this chatbot.",
        ),
        (
            "Free limits",
            "Free input/output is available. Limits are RPM, TPM and RPD per project; the exact active values are shown in AI Studio.",
            "No vendor request quota. Capacity is the server itself; busy requests queue and overload can return 503.",
            "Starter credits reset monthly, but Ollama does not publish the exact free amount. Free allows only 1 concurrent request; extra requests queue, then reject if full.",
        ),
        (
            "Privacy",
            "External. Unpaid-service terms allow submitted content to be used for product improvement. This app redacts personal data first.",
            "Strongest control: prompts, KB context and answers stay on the university-managed server.",
            "External. Ollama says prompts/responses are not logged or trained; processing is mainly in the US, with Europe/Singapore possible.",
        ),
        (
            "Real cost",
            "Starts at $0. If paid later, the measured short RAG calls are inexpensive (worked example on page 2).",
            "Model software is free, but GPU/CPU capacity, electricity, monitoring, updates and staff time are not.",
            "Starts at $0 with limited credits. Pro is $20/month for 3 concurrent requests; Max is $100 for 10; Team is $500 for 10. Extra use is token-priced.",
        ),
        (
            "Maintenance",
            "Low: API key, quota monitoring and occasional model evaluation.",
            "Higher: install and patch Ollama, store models, manage drivers/GPU, secure the service, monitor RAM/queues and recover failures.",
            "Low-to-medium: Ollama runs the compute, but model availability, API key, billing and vendor monitoring remain.",
        ),
        (
            "MSFEA fit",
            "Best fit now for the pilot; can move to paid quota if adoption grows.",
            "Suitable only if AUB provides a supported GPU server, an owner and load-test evidence.",
            "Free plan is not MSFEA-wide capacity. It also does not remove the external-provider dependency.",
        ),
    ]

    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.allow_autofit = False
    set_table_borders(table)
    widths = [1.05, 2.12, 2.12, 2.11]
    header = table.rows[0]
    set_repeat_table_header(header)
    for idx, cell in enumerate(header.cells):
        set_cell_width(cell, widths[idx])
        shade(cell, DARK)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell, 80, 70, 80, 70)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(headers[idx])
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(8.2)
        r.font.color.rgb = RGBColor(255, 255, 255)

    for row_index, values in enumerate(rows, start=1):
        row = table.add_row()
        set_cant_split(row)
        for idx, value in enumerate(values):
            cell = row.cells[idx]
            set_cell_width(cell, widths[idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, 55, 65, 55, 65)
            if row_index % 2 == 0:
                shade(cell, VERY_LIGHT)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 0.96
            r = p.add_run(value)
            r.font.name = "Arial"
            r.font.size = Pt(7.75)
            r.font.color.rgb = RGBColor.from_string(BLACK)
            if idx == 0:
                r.bold = True


def add_evidence_table(doc: Document) -> None:
    headers = ["Test case", "Evidence and result"]
    rows = [
        (
            "Normal rule application",
            "Question: 'I completed 88 credits. Can I register?' Gemini correctly concluded no because the KB minimum is 90 credits.",
        ),
        (
            "Follow-up",
            "'Can I do the ECE 6+2 arrangement?' followed by 'So the two weeks are research?' produced a direct confirmation: six company weeks plus two faculty-research weeks.",
        ),
        (
            "Unrelated second question",
            "A self-contained topic switch ignored earlier history, preventing the first topic from contaminating retrieval.",
        ),
        (
            "Uncovered or off-topic",
            "A calibrated 0.60 similarity gate refuses low-relevance questions before any LLM call and gives the CDC escalation contact.",
        ),
        (
            "Regression checks",
            "132 automated tests, Ruff and strict mypy passed. The four added reasoning cases had 4/4 retrieval context recall at top 7.",
        ),
        (
            "Local Ollama trial",
            "My informal Llama 3.1 8B test produced lower-quality answers than Gemini. No controlled accuracy or latency dataset was recorded, so this is directional evidence, not a final benchmark.",
        ),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.allow_autofit = False
    set_table_borders(table)
    widths = [1.78, 5.62]
    header = table.rows[0]
    set_repeat_table_header(header)
    for idx, cell in enumerate(header.cells):
        set_cell_width(cell, widths[idx])
        shade(cell, DARK)
        set_cell_margins(cell, 70, 80, 70, 80)
        p = cell.paragraphs[0]
        r = p.add_run(headers[idx])
        r.bold = True
        r.font.name = "Aptos"
        r.font.size = Pt(8.6)
        r.font.color.rgb = RGBColor(255, 255, 255)
    for row_index, values in enumerate(rows, start=1):
        row = table.add_row()
        set_cant_split(row)
        for idx, value in enumerate(values):
            cell = row.cells[idx]
            set_cell_width(cell, widths[idx])
            set_cell_margins(cell, 50, 75, 50, 75)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2 == 0:
                shade(cell, VERY_LIGHT)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 0.98
            r = p.add_run(value)
            r.font.name = "Aptos"
            r.font.size = Pt(8.2)
            if idx == 0:
                r.bold = True


def add_architecture_table(doc: Document) -> None:
    table = doc.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.allow_autofit = False
    set_table_borders(table)
    widths = [3.7, 3.7]
    headings = ["What stays unchanged", "What changes"]
    bodies = [
        "Source documents, ingestion, PostgreSQL/pgvector, local embeddings, hybrid retrieval, similarity refusal gate, bounded follow-up memory, grounded prompt, citations, FastAPI, frontend, dashboard and interaction metrics.",
        "Always: one LLM provider adapter and model configuration. Local Ollama also adds an Ollama service, model storage, GPU/RAM capacity, queue settings, security, updates and monitoring. Ollama Cloud adds an API key, cloud model selection, usage monitoring and billing.",
    ]
    for idx, cell in enumerate(table.rows[0].cells):
        set_cell_width(cell, widths[idx])
        shade(cell, DARK)
        set_cell_margins(cell, 70, 85, 70, 85)
        p = cell.paragraphs[0]
        r = p.add_run(headings[idx])
        r.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(255, 255, 255)
    for idx, cell in enumerate(table.rows[1].cells):
        set_cell_width(cell, widths[idx])
        set_cell_margins(cell, 70, 85, 70, 85)
        shade(cell, VERY_LIGHT)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(bodies[idx])
        r.font.size = Pt(8.6)


def build() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph()
    style_paragraph(title, after=1)
    run = title.add_run("LLM Choice for the MSFEA Chatbot")
    run.font.name = "Arial"
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(BLACK)

    subtitle = doc.add_paragraph()
    style_paragraph(subtitle, after=1)
    run = subtitle.add_run("Capstone Assignment 1 - decision brief")
    run.font.name = "Arial"
    run.font.size = Pt(10.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(ACCENT)

    meta = doc.add_paragraph()
    style_paragraph(meta, after=7)
    run = meta.add_run("Prepared by Jad Ghazi | 7 September 2026")
    run.font.size = Pt(8.7)
    run.font.color.rgb = RGBColor.from_string(MID)

    add_heading(doc, "Recommendation", 1)
    p = doc.add_paragraph()
    style_paragraph(p, after=5)
    lead = p.add_run("Keep Gemini Flash-Lite for the pilot and the capstone delivery. ")
    lead.bold = True
    p.add_run(
        "It has produced better answers in this project, responds quickly, needs no GPU, and can start at zero cost. "
        "Keep the existing provider interface so Ollama or another approved provider can be tested later without rewriting the chatbot. "
        "Do not place Llama 3.1 8B on the current CPU-only Oracle VM for student traffic."
    )

    add_heading(doc, "Side-by-side comparison", 1)
    add_comparison_table(doc)
    note = doc.add_paragraph()
    style_paragraph(note, before=3, after=0, line=1.0)
    r = note.add_run("Availability note. ")
    r.bold = True
    note.add_run(
        "Ollama's official Llama 3.1 page lists local 8B/70B/405B downloads but does not mark Llama 3.1 8B as a first-party cloud model. "
        "A cloud test may therefore require a different supported model, which is not an exact like-for-like comparison."
    )
    for run in note.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(MID)

    doc.add_page_break()

    add_heading(doc, "Evidence from this chatbot", 1)
    p = doc.add_paragraph(
        "The Gemini results below were measured in the current RAG pipeline. The Ollama observation came from an informal local trial and is reported honestly as such."
    )
    style_paragraph(p, after=4)
    add_evidence_table(doc)

    add_heading(doc, "How conversational memory works", 1)
    add_numbered(
        doc,
        "Short-lived memory. ",
        "The web page keeps at most four earlier messages in memory. A refresh, new tab or department change clears them; no permanent student profile is created.",
    )
    add_numbered(
        doc,
        "Follow-up detection. ",
        "For a short reference such as 'What are the other two weeks?', the backend combines the last two user questions with the current question for retrieval. A complete new topic ignores history.",
    )
    add_numbered(
        doc,
        "Fresh evidence. ",
        "The bot retrieves the relevant KB section, then the LLM answers from that evidence and cites it. Earlier assistant text helps resolve the reference but is never treated as a factual source.",
    )
    example = doc.add_paragraph()
    style_paragraph(example, before=2, after=4, line=1.02)
    r = example.add_run("Example: ")
    r.bold = True
    example.add_run(
        "'Can an ECE student use the 6+2 arrangement?' -> 'What are the other 2 weeks?' is retrieved as one contextual query, so the answer is 'two weeks of faculty research,' with the ECE source cited."
    )

    add_heading(doc, "Can the current Gemini setup stay free?", 1)
    add_bullet(
        doc,
        "For a small pilot, probably yes. ",
        "The model has free input/output pricing. Google applies project-level requests-per-minute, tokens-per-minute and requests-per-day limits; the exact active numbers now appear in AI Studio rather than as one guaranteed public table.",
    )
    add_bullet(
        doc,
        "The design conserves quota. ",
        "Each answered turn makes at most one Gemini call. Follow-up reformulation is local, and low-similarity/off-topic questions are refused before Gemini."
    )
    add_bullet(
        doc,
        "Measured usage is small. ",
        "Five live cases used 2,416-2,589 input tokens and 26-80 output tokens. At the current Gemini 3.5 Flash-Lite standard paid rate ($0.30/M input, $2.50/M output), a 2,600+80-token answer is about $0.00098; 10,000 similar answers are about $9.80 before taxes."
    )
    add_bullet(
        doc,
        "Operational rule. ",
        "Before inviting students, record the active AI Studio RPM/TPM/RPD values. If a 429 limit is reached, the chatbot already gives a temporary retry message instead of a wrong answer."
    )

    doc.add_page_break()

    add_heading(doc, "What actually changes if we switch?", 1)
    add_architecture_table(doc)

    add_heading(doc, "Local Ollama capacity: what 'free' really means", 1)
    p = doc.add_paragraph()
    style_paragraph(p, after=3)
    p.add_run(
        "The common Llama 3.1 8B Ollama package is about 4.9 GB, but that is only the model file. Runtime memory and the conversation context need additional RAM/VRAM, and every parallel request increases memory use."
    )
    add_bullet(
        doc,
        "Current Oracle VM (2 OCPU, 12 GB RAM, no GPU): ",
        "the model may load in a compressed form, but CPU generation would compete with the API and database. Treat it as a one-user experiment, not a dependable pilot service."
    )
    add_bullet(
        doc,
        "Practical pilot starting point: ",
        "a supported dedicated GPU with roughly 12-16 GB VRAM and at least 16 GB system RAM, then measure the real prompt, answer time and concurrent users. This is a sizing estimate, not a guarantee."
    )
    add_bullet(
        doc,
        "Concurrency and overload: ",
        "start at one generation at a time. Ollama defaults to one parallel request per model; increasing parallelism raises memory use. Busy requests queue, and a full queue returns 503."
    )
    add_bullet(
        doc,
        "MSFEA-wide operation: ",
        "requires an owned server budget, GPU monitoring, patching, backups, access control and a support person. Token fees disappear, but compute and maintenance costs replace them."
    )

    add_heading(doc, "Final decision and safeguards", 1)
    p = doc.add_paragraph()
    style_paragraph(p, after=3)
    r = p.add_run("Selected option: current external LLM (Gemini), behind the existing configurable provider interface.")
    r.bold = True
    add_bullet(doc, "Why: ", "best observed answer quality, low latency, lowest operational effort and very low fallback cost for this short-answer RAG workload.")
    add_bullet(doc, "Before launch: ", "confirm the project's live free-tier limits, keep personal-data redaction, monitor 429s/tokens/latency, and pin a stable Gemini model after re-running the golden set so a moving 'latest' alias cannot change behavior silently.")
    add_bullet(doc, "If local hosting is reconsidered: ", "obtain real AUB GPU hardware, run the same golden answer set against both models, and load-test concurrent student traffic before changing the provider.")
    add_bullet(doc, "No automatic hybrid yet: ", "automatic failover adds privacy, consistency and testing complexity. A manual provider switch is enough for the pilot."
    )

    add_heading(doc, "Sources checked", 2)
    sources = [
        ("Google Gemini API pricing", "https://ai.google.dev/gemini-api/docs/pricing"),
        ("Google Gemini API rate limits", "https://ai.google.dev/gemini-api/docs/rate-limits"),
        ("Google Gemini API terms", "https://ai.google.dev/gemini-api/terms"),
        ("Ollama pricing and cloud concurrency", "https://ollama.com/pricing"),
        ("Ollama Llama 3.1 model page", "https://ollama.com/library/llama3.1"),
        ("Ollama concurrency and overload FAQ", "https://docs.ollama.com/faq"),
    ]
    p = doc.add_paragraph()
    style_paragraph(p, after=1, line=1.0)
    for index, (label, url) in enumerate(sources, start=1):
        p.add_run(f"{index}. ").font.size = Pt(7.8)
        add_hyperlink(p, label, url)
        if index < len(sources):
            p.add_run("  |  ")
    for run in p.runs:
        run.font.name = "Arial"
        run.font.size = Pt(7.8)
    p2 = doc.add_paragraph(
        "Project evidence: docs/decisions/0018, docs/decisions/0021 and docs/progress.md in the MSFEA Chatbot repository. Web facts checked 7 September 2026; vendor plans and quotas can change."
    )
    style_paragraph(p2, after=0, line=1.0)
    for run in p2.runs:
        run.font.size = Pt(7.6)
        run.font.color.rgb = RGBColor.from_string(MID)

    core = doc.core_properties
    core.title = "LLM Choice for the MSFEA Chatbot"
    core.subject = "Capstone Assignment 1 - comparison of Gemini, Ollama local, and Ollama Cloud"
    core.author = "Jad Ghazi"
    core.keywords = "MSFEA, chatbot, RAG, Gemini, Ollama, Llama 3.1"
    core.comments = "Prepared from project measurements and official vendor documentation."

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
